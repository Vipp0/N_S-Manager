"""Export dello stato corrente del database verso un file Excel a 3 fogli."""
import sqlite3
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from gilda_app.db.database import get_members
from gilda_app.models.member import STATUS_ATTIVO, STATUS_BANNATO, STATUS_EX_MEMBRO

HEADERS = ["Family Name", "Main Name", "Nation", "Discord Name", "Note"]

SHEET_FOR_STATUS = [
    ("Members", STATUS_ATTIVO),
    ("Old Members", STATUS_EX_MEMBRO),
    ("No Rejoin", STATUS_BANNATO),
]


def export_workbook(conn: sqlite3.Connection, path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name, status in SHEET_FOR_STATUS:
        ws = wb.create_sheet(sheet_name)
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for member in get_members(conn, status):
            ws.append(
                [
                    member.family_name,
                    member.main_name,
                    "/".join(member.nations),
                    member.discord_name,
                    member.note or "",
                ]
            )
        for col_idx in range(1, len(HEADERS) + 1):
            ws.column_dimensions[chr(ord("A") + col_idx - 1)].width = 22

    wb.save(str(path))
